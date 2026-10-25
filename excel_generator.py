"""
AWS Infrastructure Specification Excel Generator.
고객사 제출용 고품격 디자인 템플릿(세련된 헤더 컬러, 카드형 요약 표, 자동 열 너비, 상태값 뱃지 서식)을 적용합니다.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 스타일 상수 정의
FONT_TITLE = Font(name="맑은 고딕", size=16, bold=True, color="1E293B")
FONT_SUBTITLE = Font(name="맑은 고딕", size=11, color="64748B")
FONT_SECTION = Font(name="맑은 고딕", size=12, bold=True, color="0F172A")
FONT_HEADER = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="맑은 고딕", size=10, color="1E293B")
FONT_BOLD = Font(name="맑은 고딕", size=10, bold=True, color="1E293B")

# 배경색 Fill
FILL_AWS_NAVY = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # 짙은 슬레이트 네이비
FILL_HEADER_ALT = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # 액센트 블루
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")      # 교차행 연그레이
FILL_CARD_BG = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")    # 카드 배경
FILL_HIGHLIGHT = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # 강조 옐로우

# 상태값 뱃지 스타일
FILL_STATE_RUNNING = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FONT_STATE_RUNNING = Font(name="맑은 고딕", size=10, bold=True, color="15803D")

FILL_STATE_STOPPED = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
FONT_STATE_STOPPED = Font(name="맑은 고딕", size=10, bold=True, color="B91C1C")

# 테두리 Border
THIN_GRAY = Side(style="thin", color="CBD5E1")
BORDER_CELL = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
BORDER_TOP_THICK = Border(top=Side(style="medium", color="1E3A8A"))
BORDER_BOTTOM_DOUBLE = Border(bottom=Side(style="double", color="1E3A8A"))


def auto_fit_columns(ws, min_width=12, max_width=55):
    """모든 열의 글자 길이를 계산하여 열 너비를 최적화"""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # 병합된 셀 제외
            if cell.coordinate in ws.merged_cells:
                continue
            val = str(cell.value or "")
            if "\n" in val:
                val = max(val.split("\n"), key=len)
            
            # 한글/영문 길이 보정 (한글은 1.8배 가중치)
            byte_len = 0
            for ch in val:
                byte_len += 1.8 if ord(ch) > 127 else 1.0
            if byte_len > max_len:
                max_len = int(byte_len)

        target_width = max(min_width, min(max_len + 4, max_width))
        ws.column_dimensions[col_letter].width = target_width


def style_header_row(ws, row_idx, headers, fill=FILL_AWS_NAVY):
    """테이블 헤더 행 서식 적용"""
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header_text)
        cell.font = FONT_HEADER
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_CELL
    ws.row_dimensions[row_idx].height = 26


def apply_data_row_style(cell, is_even=False, align="left"):
    """일반 데이터 셀 서식 적용"""
    cell.font = FONT_DATA
    cell.border = BORDER_CELL
    if is_even:
        cell.fill = FILL_ZEBRA
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


class ExcelReportBuilder:
    def __init__(self, data: dict, customer_name: str = "고객사"):
        self.data = data
        self.customer = customer_name or "고객사"
        self.region = data.get("metadata", {}).get("region", "ap-northeast-2")
        self.scanned_at = data.get("metadata", {}).get("scanned_at", "-")
        self.summary = data.get("summary", {})
        self.wb = Workbook()
        # 기본 시트 제거
        self.wb.remove(self.wb.active)

    def generate(self) -> Workbook:
        """전체 엑셀 시트 생성"""
        self._build_overview_sheet()
        self._build_network_sheet()
        self._build_gateways_routes_sheet()
        self._build_security_groups_sheet()
        self._build_ec2_sheet()
        self._build_rds_sheet()
        self._build_elb_sheet()
        self._build_s3_sheet()
        self._build_lambda_sheet()
        return self.wb

    def save_to_file(self, filename: str):
        self.generate()
        self.wb.save(filename)

    def save_to_bytes(self) -> io.BytesIO:
        self.generate()
        output = io.BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output

    def _build_overview_sheet(self):
        ws = self.wb.create_sheet(title="00_인프라_요약_보고서")
        ws.views.sheetView[0].showGridLines = True

        # 메인 타이틀
        ws.merge_cells("B2:G2")
        title_cell = ws["B2"]
        title_cell.value = "AWS 클라우드 인프라 생성 명세서"
        title_cell.font = FONT_TITLE
        title_cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 32

        # 서브 헤더
        ws.merge_cells("B3:G3")
        sub_cell = ws["B3"]
        sub_cell.value = f"본 문서는 {self.customer} 제출용으로 자동 추출된 AWS 핵심 코어 및 네트워크 자원 구성 명세서입니다."
        sub_cell.font = FONT_SUBTITLE
        ws.row_dimensions[3].height = 20

        # 프로젝트 메타데이터 테이블 (카드 형태)
        meta_data = [
            ("고객사명", self.customer, "기준 리전 (Region)", f"{self.region}"),
            ("문서 유형", "AWS 인프라 생성 명세서 (최종)", "생성 일시 (Scanned At)", self.scanned_at),
        ]
        curr_row = 5
        for row in meta_data:
            ws.cell(row=curr_row, column=2, value=row[0]).font = FONT_BOLD
            ws.cell(row=curr_row, column=2).fill = FILL_CARD_BG
            ws.cell(row=curr_row, column=2).border = BORDER_CELL
            ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center", vertical="center")

            c1 = ws.cell(row=curr_row, column=3, value=row[1])
            c1.font = FONT_DATA
            c1.border = BORDER_CELL
            c1.alignment = Alignment(vertical="center")

            ws.cell(row=curr_row, column=4, value=row[2]).font = FONT_BOLD
            ws.cell(row=curr_row, column=4).fill = FILL_CARD_BG
            ws.cell(row=curr_row, column=4).border = BORDER_CELL
            ws.cell(row=curr_row, column=4).alignment = Alignment(horizontal="center", vertical="center")

            c2 = ws.cell(row=curr_row, column=5, value=row[3])
            c2.font = FONT_DATA
            c2.border = BORDER_CELL
            c2.alignment = Alignment(vertical="center")
            
            ws.row_dimensions[curr_row].height = 24
            curr_row += 1

        # 섹션 헤더: 핵심 자원 보유 현황 KPI 요약
        curr_row = 8
        ws.cell(row=curr_row, column=2, value="■ 핵심 인프라 자원 수량 요약").font = FONT_SECTION
        curr_row += 1

        kpi_headers = ["No", "서비스 영역", "자원 구분 (Resource)", "인프라 수량", "핵심 설명"]
        for idx, h in enumerate(kpi_headers, 2):
            c = ws.cell(row=curr_row, column=idx, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_AWS_NAVY
            c.border = BORDER_CELL
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[curr_row].height = 26
        curr_row += 1

        kpis = [
            (1, "네트워크 (Network)", "VPC (가상 프라이빗 클라우드)", f"{self.summary.get('vpc_count', 0)} 개", "격리된 가상 네트워크 환경"),
            (2, "네트워크 (Network)", "서브넷 (Subnet)", f"{self.summary.get('subnet_count', 0)} 개", "Public / Private 망 분리 서브넷"),
            (3, "네트워크 (Network)", "인터넷 / NAT 게이트웨이", f"{self.summary.get('igw_count', 0)} / {self.summary.get('nat_count', 0)} 개", "외부 인터넷 및 아웃바운드 게이트웨이"),
            (4, "보안 (Security)", "보안 그룹 (Security Group)", f"{self.summary.get('sg_count', 0)} 개", "가상 방화벽 인바운드/아웃바운드 제어"),
            (5, "컴퓨팅 (Compute)", "EC2 인스턴스", f"{self.summary.get('ec2_count', 0)} 대", "애플리케이션 및 WAS 가상 서버"),
            (6, "데이터베이스 (DB)", "RDS / Aurora 인스턴스", f"{self.summary.get('rds_count', 0)} 개", "고가용성 관리형 관계형 데이터베이스"),
            (7, "로드밸런서 (ELB)", "Application / Network LB", f"{self.summary.get('elb_count', 0)} 대", "서비스 트래픽 부하분산 및 리스너"),
            (8, "스토리지 (Storage)", "S3 버킷", f"{self.summary.get('s3_count', 0)} 개", "정적 웹 호스팅 및 백업/오브젝트 저장소"),
            (9, "서버리스 (Serverless)", "Lambda 함수", f"{self.summary.get('lambda_count', 0)} 개", "이벤트 기반 서버리스 컴퓨팅 로직"),
        ]

        for item in kpis:
            is_even = item[0] % 2 == 0
            for c_idx, val in enumerate(item, 2):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                c.font = FONT_DATA
                c.border = BORDER_CELL
                if is_even:
                    c.fill = FILL_ZEBRA
                if c_idx in [2, 5]: # No, 수량
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[curr_row].height = 22
            curr_row += 1

        auto_fit_columns(ws)
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 30
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 42

    def _build_network_sheet(self):
        ws = self.wb.create_sheet(title="01_네트워크_VPC_서브넷")
        ws.views.sheetView[0].showGridLines = True

        # 1. VPC 섹션
        ws.cell(row=1, column=1, value="■ VPC (Virtual Private Cloud) 명세").font = FONT_SECTION
        vpc_headers = ["Name (태그)", "VPC ID", "IPv4 CIDR 블록", "구분 (Default 여부)", "상태 (State)", "DHCP 옵션 ID"]
        style_header_row(ws, 2, vpc_headers)

        curr_row = 3
        vpcs = self.data.get("vpcs", [])
        if not vpcs:
            ws.cell(row=curr_row, column=1, value="조회된 VPC 자원이 없습니다.").font = FONT_DATA
            curr_row += 1
        else:
            for idx, v in enumerate(vpcs):
                is_even = idx % 2 == 1
                row_vals = [v["name"], v["vpc_id"], v["cidr"], v["is_default"], v["state"], v["dhcp"]]
                for c_idx, val in enumerate(row_vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 5] else "left")
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        curr_row += 2
        # 2. Subnet 섹션
        ws.cell(row=curr_row, column=1, value="■ 서브넷 (Subnet) 상세 명세").font = FONT_SECTION
        curr_row += 1
        sub_headers = ["서브넷 이름", "서브넷 ID", "소속 VPC ID", "IPv4 CIDR", "가용영역 (AZ)", "망 구분 (Type)", "가용 IP 수", "상태"]
        style_header_row(ws, curr_row, sub_headers, fill=FILL_HEADER_ALT)
        curr_row += 1

        subnets = self.data.get("subnets", [])
        if not subnets:
            ws.cell(row=curr_row, column=1, value="조회된 서브넷 자원이 없습니다.").font = FONT_DATA
        else:
            for idx, s in enumerate(subnets):
                is_even = idx % 2 == 1
                row_vals = [s["name"], s["subnet_id"], s["vpc_id"], s["cidr"], s["az"], s["type"], s["available_ips"], s["state"]]
                for c_idx, val in enumerate(row_vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 5, 6, 8] else ("right" if c_idx == 7 else "left"))
                    if c_idx == 6: # Type (Public/Private)
                        if s["type"] == "Public":
                            c.font = Font(name="맑은 고딕", size=10, bold=True, color="2563EB")
                        else:
                            c.font = Font(name="맑은 고딕", size=10, bold=True, color="475569")
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        auto_fit_columns(ws)

    def _build_gateways_routes_sheet(self):
        ws = self.wb.create_sheet(title="02_게이트웨이_라우팅")
        ws.views.sheetView[0].showGridLines = True

        # IGW
        ws.cell(row=1, column=1, value="■ 인터넷 게이트웨이 (Internet Gateway)").font = FONT_SECTION
        igw_headers = ["IGW 이름", "인터넷 게이트웨이 ID", "연결된 VPC ID", "연결 상태"]
        style_header_row(ws, 2, igw_headers)
        curr_row = 3
        igws = self.data.get("gateways", {}).get("igw", [])
        for idx, g in enumerate(igws):
            is_even = idx % 2 == 1
            for c_idx, val in enumerate([g["name"], g["igw_id"], g["attached_vpc"], g["state"]], 1):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4] else "left")
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        curr_row += 2
        # NAT GW
        ws.cell(row=curr_row, column=1, value="■ NAT 게이트웨이 (NAT Gateway)").font = FONT_SECTION
        curr_row += 1
        nat_headers = ["NAT GW 이름", "NAT 게이트웨이 ID", "소속 VPC ID", "배치 서브넷 ID", "탄력적 IP (Elastic IP)", "상태"]
        style_header_row(ws, curr_row, nat_headers, fill=FILL_HEADER_ALT)
        curr_row += 1
        nats = self.data.get("gateways", {}).get("nat", [])
        for idx, n in enumerate(nats):
            is_even = idx % 2 == 1
            for c_idx, val in enumerate([n["name"], n["nat_id"], n["vpc_id"], n["subnet_id"], n["public_ip"], n["state"]], 1):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 5, 6] else "left")
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        curr_row += 2
        # Route Table
        ws.cell(row=curr_row, column=1, value="■ 라우팅 테이블 (Route Table)").font = FONT_SECTION
        curr_row += 1
        rtb_headers = ["라우팅 테이블 이름", "Route Table ID", "소속 VPC ID", "연결된 서브넷", "기본 게이트웨이 (0.0.0.0/0)", "라우팅 규칙 요약"]
        style_header_row(ws, curr_row, rtb_headers)
        curr_row += 1
        rtbs = self.data.get("route_tables", [])
        for idx, r in enumerate(rtbs):
            is_even = idx % 2 == 1
            for c_idx, val in enumerate([r["name"], r["rtb_id"], r["vpc_id"], r["associated_subnets"], r["default_gateway"], r["routes_summary"]], 1):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 5] else "left")
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        auto_fit_columns(ws)

    def _build_security_groups_sheet(self):
        ws = self.wb.create_sheet(title="03_보안그룹")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ 보안 그룹 (Security Group) 인바운드 방화벽 규칙 명세").font = FONT_SECTION
        headers = ["보안그룹 명칭 (Name)", "보안그룹 ID", "그룹 코드명", "소속 VPC ID", "인바운드 허용 룰 요약", "설명 (Description)"]
        style_header_row(ws, 2, headers)

        curr_row = 3
        sgs = self.data.get("security_groups", [])
        for idx, sg in enumerate(sgs):
            is_even = idx % 2 == 1
            vals = [sg["name"], sg["group_id"], sg["group_name"], sg["vpc_id"], sg["inbound_rules"], sg["description"]]
            for c_idx, val in enumerate(vals, 1):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                apply_data_row_style(c, is_even, align="center" if c_idx in [2, 4] else "left")
            ws.row_dimensions[curr_row].height = 22 if "\n" not in sg["inbound_rules"] else 36
            curr_row += 1

        auto_fit_columns(ws)

    def _build_ec2_sheet(self):
        ws = self.wb.create_sheet(title="04_컴퓨팅_EC2")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ EC2 가상 서버 인스턴스 명세").font = FONT_SECTION
        headers = [
            "인스턴스 이름 (Name)", "인스턴스 ID", "상태 (State)", "인스턴스 타입",
            "vCPU", "메모리(GiB)", "Private IP", "Public IP", "소속 서브넷 ID", "가용영역 (AZ)", "Key Pair", "생성/시작일시"
        ]
        style_header_row(ws, 2, headers)

        curr_row = 3
        instances = self.data.get("ec2_instances", [])
        if not instances:
            ws.cell(row=curr_row, column=1, value="조회된 EC2 인스턴스가 없습니다.").font = FONT_DATA
        else:
            for idx, inst in enumerate(instances):
                is_even = idx % 2 == 1
                vals = [
                    inst["name"], inst["instance_id"], inst["state"], inst["instance_type"],
                    inst["vcpu"], inst["ram_gib"], inst["private_ip"], inst["public_ip"],
                    inst["subnet_id"], inst["az"], inst["key_name"], inst["launch_time"]
                ]
                for c_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12] else "left")
                    if c_idx == 3: # State badge
                        if inst["state"] == "running":
                            c.fill = FILL_STATE_RUNNING
                            c.font = FONT_STATE_RUNNING
                        elif inst["state"] == "stopped":
                            c.fill = FILL_STATE_STOPPED
                            c.font = FONT_STATE_STOPPED
                ws.row_dimensions[curr_row].height = 22
                curr_row += 1

        auto_fit_columns(ws)

    def _build_rds_sheet(self):
        ws = self.wb.create_sheet(title="05_데이터베이스_RDS")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ 관계형 데이터베이스 (RDS / Aurora) 명세").font = FONT_SECTION
        headers = [
            "DB 식별자 (Identifier)", "상태 (Status)", "엔진 및 버전", "인스턴스 클래스",
            "스토리지 용량", "Multi-AZ 여부", "엔드포인트 접속 주소", "소속 VPC ID", "DB 서브넷 그룹"
        ]
        style_header_row(ws, 2, headers)

        curr_row = 3
        dbs = self.data.get("rds_instances", [])
        if not dbs:
            ws.cell(row=curr_row, column=1, value="조회된 RDS 인스턴스가 없습니다.").font = FONT_DATA
        else:
            for idx, db in enumerate(dbs):
                is_even = idx % 2 == 1
                vals = [
                    db["db_id"], db["status"], db["engine"], db["db_class"],
                    db["storage"], db["multi_az"], db["endpoint"], db["vpc_id"], db["subnet_group"]
                ]
                for c_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 4, 5, 6, 8] else "left")
                    if c_idx == 2 and db["status"] == "available":
                        c.fill = FILL_STATE_RUNNING
                        c.font = FONT_STATE_RUNNING
                ws.row_dimensions[curr_row].height = 22
                curr_row += 1

        auto_fit_columns(ws)

    def _build_elb_sheet(self):
        ws = self.wb.create_sheet(title="06_로드밸런서_ELB")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ Application / Network Load Balancer (ELB) 명세").font = FONT_SECTION
        headers = ["로드밸런서 이름", "유형 (Type)", "스킴 (Scheme)", "상태", "DNS 엔드포인트 주소", "소속 VPC ID", "리스너 포트"]
        style_header_row(ws, 2, headers)

        curr_row = 3
        elbs = self.data.get("load_balancers", [])
        if not elbs:
            ws.cell(row=curr_row, column=1, value="조회된 로드밸런서가 없습니다.").font = FONT_DATA
        else:
            for idx, lb in enumerate(elbs):
                is_even = idx % 2 == 1
                vals = [lb["name"], lb["type"], lb["scheme"], lb["state"], lb["dns_name"], lb["vpc_id"], lb["listeners"]]
                for c_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 6, 7] else "left")
                    if c_idx == 4 and lb["state"] == "active":
                        c.fill = FILL_STATE_RUNNING
                        c.font = FONT_STATE_RUNNING
                ws.row_dimensions[curr_row].height = 22
                curr_row += 1

        auto_fit_columns(ws)

    def _build_s3_sheet(self):
        ws = self.wb.create_sheet(title="07_스토리지_S3")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ Simple Storage Service (S3) 버킷 명세").font = FONT_SECTION
        headers = ["버킷 이름 (Bucket Name)", "배치 리전 (Region)", "현재 기준 리전 여부", "생성 일시"]
        style_header_row(ws, 2, headers)

        curr_row = 3
        s3s = self.data.get("s3_buckets", [])
        if not s3s:
            ws.cell(row=curr_row, column=1, value="조회된 S3 버킷이 없습니다.").font = FONT_DATA
        else:
            for idx, b in enumerate(s3s):
                is_even = idx % 2 == 1
                vals = [b["name"], b["region"], b["is_current_region"], b["created_at"]]
                for c_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4] else "left")
                    if c_idx == 3 and b["is_current_region"] == "예":
                        c.font = Font(name="맑은 고딕", size=10, bold=True, color="15803D")
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        auto_fit_columns(ws)

    def _build_lambda_sheet(self):
        ws = self.wb.create_sheet(title="08_서버리스_Lambda")
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=1, column=1, value="■ AWS Lambda 함수 명세").font = FONT_SECTION
        headers = ["함수 이름 (Function Name)", "런타임 (Runtime)", "할당 메모리", "타임아웃", "핸들러 (Handler)", "최근 수정일시"]
        style_header_row(ws, 2, headers)

        curr_row = 3
        funcs = self.data.get("lambda_functions", [])
        if not funcs:
            ws.cell(row=curr_row, column=1, value="조회된 Lambda 함수가 없습니다.").font = FONT_DATA
        else:
            for idx, f in enumerate(funcs):
                is_even = idx % 2 == 1
                vals = [f["name"], f["runtime"], f["memory"], f["timeout"], f["handler"], f["last_modified"]]
                for c_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=curr_row, column=c_idx, value=val)
                    apply_data_row_style(c, is_even, align="center" if c_idx in [2, 3, 4, 6] else "left")
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        auto_fit_columns(ws)
